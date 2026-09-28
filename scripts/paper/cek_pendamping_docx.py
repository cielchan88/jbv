"""Periksa angka naskah pendamping yang SUDAH dibangun, terhadap tabelnya sendiri.

    python3 scripts/paper/cek_pendamping_docx.py <docx-hasil-build>

BEDANYA DENGAN cek_naskah_pendamping.py. Yang itu mengaudit SUMBER build_paper.js
secara statis - aritmetika angka yang diketik, penomoran TMAP/T, angka yang
dipakai ulang. Yang ini membaca DOCX HASILNYA dan mencocokkan prosa dengan sel
tabel, yang baru mungkin setelah naskah bisa disusun ulang.

Ini sepadan dengan cek_draft.py milik naskah utama, dengan satu beda yang harus
dinyatakan: naskah utama mencocokkan angka dengan HASIL KOMPUTASI, sedangkan di
sini tidak ada hasil komputasi di repo. Yang bisa diperiksa adalah apakah naskah
KONSISTEN DENGAN DIRINYA SENDIRI. Itu menangkap angka usang dan salah rujuk;
itu TIDAK menangkap tabel yang seluruhnya salah.
"""
import os
import re
import sys
from collections import Counter

import docx

gagal, catatan = [], []
fail = lambda l, m: gagal.append(f'{l}: {m}')


def sel(t):
    return [[c.text.strip() for c in r.cells] for r in t.rows]


def angka(s):
    s = s.replace('−', '-').replace('×', '').replace('%', '').strip()
    m = re.match(r'^[+-]?[\d.]+$', s)
    return float(s) if m else None


def main():
    if len(sys.argv) < 2:
        raise SystemExit('Pemakaian: cek_pendamping_docx.py <docx-hasil-build>')
    F = sys.argv[1]
    if not os.path.exists(F):
        raise SystemExit(f'BERHENTI: {F} tidak ada.')
    d = docx.Document(F)
    prosa = '\n'.join(p.text for p in d.paragraphs)
    tabel = [sel(t) for t in d.tables]

    # Keterangan -> tabel, dibaca menurut urutan badan dokumen supaya pasangannya
    # benar-benar yang dicetak, bukan tebakan dari indeks.
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    ket, pasang, nomor = None, {}, []
    for el in d.element.body.iterchildren():
        tag = el.tag.split('}')[1]
        if tag == 'p':
            t = Paragraph(el, d).text.strip()
            m = re.match(r'^Table ([A-Z]?\d+)\.', t)
            if m:
                ket = m.group(1)
                nomor.append(ket)
        elif tag == 'tbl':
            if ket:
                pasang[ket] = sel(Table(el, d))
            ket = None

    print('1. penomoran dan rujukan')
    ganda = {k: v for k, v in Counter(nomor).items() if v > 1}
    if ganda:
        fail('keterangan', f'nomor tabel ganda: {ganda}')
    tanpa = [n for n in nomor if n not in pasang]
    if tanpa:
        fail('keterangan', f'keterangan tanpa tabel sesudahnya: {tanpa}')
    print(f'   {len(pasang)} tabel berketerangan: {", ".join(nomor)}')

    dirujuk = {m for m in re.findall(r'Tabl?e ([A-Z]?\d+)', prosa)}
    hilang = sorted(dirujuk - set(pasang))
    if hilang:
        fail('rujukan', f'prosa menyebut Table {hilang} yang tidak ada')
    gbr_ada = {m for m in re.findall(r'^Figure ([A-Z]?\d+)\.', prosa, re.M)}
    gbr_rujuk = {m for m in re.findall(r'Figure ([A-Z]?\d+)', prosa)}
    hilang = sorted(gbr_rujuk - gbr_ada)
    if hilang:
        fail('rujukan', f'prosa menyebut Figure {hilang} yang tidak punya keterangan')
    print(f'   gambar berketerangan: {", ".join(sorted(gbr_ada))}')

    # ---------------------------------------------------- 2 prosa lawan tabel
    print('\n2. prosa lawan sel tabel')

    def cari_tabel(kepala):
        for n, t in pasang.items():
            if t and t[0][:len(kepala)] == kepala:
                return n, t
        return None, None

    # Tahap I: klaim abstrak menyebut nilai yang harus ada di tabelnya.
    n6, t6 = cari_tabel(['Method', 'Family', 'In MASE'])
    if not t6:
        fail('Tahap I', 'tabel fit dalam-sampel tidak ketemu')
    else:
        d6 = {r[0]: r for r in t6[1:]}
        for metode, kol, label in (('RandomForest', 2, 'in-sample'),
                                   ('RandomForest', 5, 'out-of-sample'),
                                   ('Prophet', 2, 'in-sample'),
                                   ('Prophet', 5, 'out-of-sample')):
            v = d6[metode][kol]
            if v not in prosa:
                fail('Tahap I', f'{metode} {label} {v} ada di Table {n6} '
                                'tapi tidak muncul di prosa')
        print(f'   Table {n6}: keempat nilai kunci muncul di prosa')
        # Pembalikan peringkat: baris tabel terurut menurut fit dalam-sampel,
        # jadi kolom Rank (out) tanpa random walk harus persis terbalik.
        tanpa_rw = [r for r in t6[1:] if r[0] != 'Naive']
        keluar = [int(r[6]) for r in tanpa_rw]
        harus = sorted(keluar, reverse=True)
        if keluar != harus:
            fail('Tahap I', f'tanpa random walk peringkat luar-sampel {keluar} '
                            'BUKAN kebalikan sempurna, padahal naskah mengklaim rho = -1')
        else:
            print(f'   tanpa random walk, {len(tanpa_rw)} metode terbalik sempurna '
                  f'-> rho = -1 sah')

    # Protokol: baris Mean harus cocok dengan prosa, dan persentasenya turunan.
    n9, t9 = cari_tabel(['Method', 'Teacher-forced'])
    if t9:
        mean = next((r for r in t9[1:] if r[0].lower() == 'mean'), None)
        if mean:
            tf, dr, rc = (angka(mean[1]), angka(mean[2]), angka(mean[3]))
            for v in (mean[1], mean[2], mean[3]):
                if v not in prosa:
                    fail('protokol', f'nilai {v} di baris Mean Table {n9} '
                                     'tidak muncul di prosa')
            pen = 100 * (rc - tf) / tf
            akum = 100 * (rc - dr) / dr
            bocor = 100 * (dr - tf) / tf
            print(f'   Table {n9} Mean: {tf}/{dr}/{rc} -> total {pen:.1f}%, '
                  f'akumulasi {akum:.1f}%, kebocoran {bocor:.1f}%')
            for nm, v in (('total', pen), ('akumulasi', akum), ('kebocoran', bocor)):
                if f'{v:.1f} per cent' not in prosa:
                    fail(f'protokol/{nm}', f'{v:.1f} per cent tidak muncul di prosa')
            if abs((akum + bocor) - pen) > 0.05 and 'percentage points' in prosa:
                fail('protokol', 'prosa memakai "percentage points" padahal '
                                 'komponennya berlipat, bukan menjumlah')

    # Segmen horizon: persentase di prosa harus lahir dari barisnya.
    n10, t10 = cari_tabel(['Horizon segment'])
    if t10:
        for r in t10[1:]:
            rc, dr = angka(r[1]), angka(r[2])
            if rc is None or dr is None:
                continue
            p = 100 * (rc - dr) / dr
            pola = rf'{r[1]} against {r[2]}, a penalty of (\d+) per cent'
            m = re.search(pola, prosa)
            if m:
                print(f'   Table {n10} {r[0]}: {rc} vs {dr} = {p:.1f}%, '
                      f'prosa {m.group(1)}%')
                if abs(p - int(m.group(1))) > 0.5:
                    fail('horizon', f'{rc} vs {dr} memberi {p:.1f}%, '
                                    f'prosa menulis {m.group(1)}%')

    # Tabel robustness memuat baris yang menyebut n; abstrak harus setuju.
    n13, t13 = cari_tabel(['Test', 'Statistic'])
    if t13:
        for r in t13[1:]:
            m = re.search(r'excluding the random walk, n = (\d+)', r[0])
            if m:
                n = int(m.group(1))
                EJA = {5: 'five', 6: 'six', 7: 'seven', 8: 'eight'}
                if not re.search(rf'\b(?:{n}|{EJA.get(n, n)})\b[^.]{{0,60}}methods',
                                 prosa, re.I):
                    fail('robustness', f'Table {n13} menyebut n = {n} tanpa random '
                                       'walk, tapi prosa tidak pernah menyebut '
                                       'jumlah metode itu')
                else:
                    print(f'   Table {n13}: n = {n} tanpa random walk, dan prosa '
                          'menyebutnya')
                if r[1] not in prosa:
                    fail('robustness', f'statistik {r[1]} tidak muncul di prosa')

    # ------------------------------------------------- 3 konsistensi antar-tabel
    print('\n3. konsistensi antar-tabel')
    n11, t11 = cari_tabel(['Method', 'Family', 'MASE', 'MAE'])
    n12, t12 = cari_tabel(['Method', 'Family', 'MASE', 'Median'])
    if t11 and t12:
        urut = {r[0]: i for i, r in enumerate(t11[1:], 1)}
        h = t12[0]
        if 'Rank h=60' in h:
            k = h.index('Rank h=60')
            salah = [(r[0], r[k], urut.get(r[0])) for r in t12[1:]
                     if str(urut.get(r[0])) != r[k]]
            if salah:
                fail('peringkat', f'kolom Rank h=60 tidak cocok dengan urutan '
                                  f'Table {n11}: {salah}')
            else:
                print(f'   kolom Rank h=60 Table {n12} cocok dengan urutan Table {n11}')
        if 'Best in' in h:
            k = h.index('Best in')
            jml = sum(int(r[k]) for r in t12[1:])
            m = re.search(r'sums to (\d+)', prosa)
            print(f'   "Best in" berjumlah {jml}' +
                  (f', prosa menyatakan {m.group(1)}' if m else ''))
            if m and jml != int(m.group(1)):
                fail('best in', f'berjumlah {jml}, prosa menyatakan {m.group(1)}')
        # Kosakata keluarga harus sama di seluruh tabel.
        kel = {}
        for n, t in pasang.items():
            if t and 'Family' in t[0]:
                k = t[0].index('Family')
                kel[n] = {r[k] for r in t[1:] if len(r) > k}
        semua = set().union(*kel.values()) if kel else set()
        if len({v for v in semua if v.lower().startswith(('ml', 'machine'))}) > 1:
            catatan.append('label keluarga tidak seragam antar-tabel: '
                           + '; '.join(f'Table {n}: {sorted(v)}' for n, v in kel.items()))

    # Rasio di tabel keluarga fitur harus bisa dilahirkan kolomnya sendiri.
    n8, t8 = cari_tabel(['Feature family'])
    if t8:
        meleset = []
        for r in t8[1:]:
            a, b, c = angka(r[1]), angka(r[2]), angka(r[3])
            if None in (a, b, c):
                continue
            if abs(b / a - c) > 0.015:
                meleset.append(f'{r[0]}: {b}/{a} = {b/a:.3f}, tercetak {c}')
        if meleset:
            catatan.append(f'Table {n8}: rasio tidak persis lahir dari kolom '
                           'pangsanya (pangsanya dibulatkan satu desimal) - '
                           + '; '.join(meleset))

    # --------------------------------------------------------- 4 nama metode
    print('\n4. nama metode')
    n_auto = prosa.count('AutoARIMA') + sum(
        c.count('AutoARIMA') for t in tabel for r in t for c in r)
    print(f'   kemunculan "AutoARIMA": {n_auto}')
    if n_auto:
        fail('nama', f'{n_auto} kemunculan "AutoARIMA" tersisa; naskah utama '
                     'menulis ARIMA')

    # ------------------------------------------------------------- laporan
    print()
    if gagal:
        print(f'GAGAL: {len(gagal)} masalah')
        for g in gagal:
            print(f'  - {g}')
    else:
        print('konsistensi internal naskah pendamping: nol masalah')
    if catatan:
        print(f'\nperlu diperiksa manusia ({len(catatan)}):')
        for c in catatan:
            print(f'  - {c}')
    print('\nCATATAN: yang diperiksa adalah konsistensi naskah DENGAN DIRINYA '
          'SENDIRI.\nTidak ada hasil komputasi naskah ini di repo, jadi tabel '
          'yang seluruhnya\nsalah tidak akan tertangkap di sini.')
    sys.exit(1 if gagal else 0)


if __name__ == '__main__':
    main()
