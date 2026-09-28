"""Pulihkan h1_facts2.json dari DOCX naskah pendamping yang SUDAH pernah dibangun.

    python3 scripts/paper/pulihkan_h1_facts.py <docx-hasil-build> -o <dir-kerja>

KENAPA SKRIP INI ADA. build_paper.js membaca h1_facts2.json, tetapi TIDAK ADA
apa pun di repo yang menulisnya: run_one_step.py menulis h1_facts.json dengan
skema berbeda ('table' bukan 'rows', 'median' bukan 'med', 'lag_based' bukan
'family', dan tanpa r1/r60/best, rho, rho_p, proto_mase, rec_tf_*, dir_*).
Angka "2" pada namanya menyiratkan langkah kedua yang tidak ikut di-commit.
Akibatnya naskah pendamping tidak bisa disusun ulang sama sekali.

Skrip ini menutup celah itu dari arah lain: seluruh isi H1 sudah TERCETAK di
DOCX hasil build sebelumnya - tabel satu langkah memuat rows, dan prosanya
memuat setiap skalar - jadi masukannya bisa dibaca kembali dari keluarannya.

BATASNYA HARUS DINYATAKAN. Yang dipulihkan adalah nilai yang SUDAH DIBULATKAN
saat dicetak: mase tiga desimal, mae satu desimal, dan seterusnya. Untuk
menyusun ulang naskah itu cukup, karena build_paper.js memang mencetaknya
dengan pembulatan yang sama. Ia BUKAN pengganti menjalankan ulang
run_one_step.py, dan tidak boleh dipakai sebagai sumber angka baru.

PENJAGA. Setiap medan wajib harus ketemu; kalau kalimat di naskah diubah
sehingga polanya meleset, skrip berhenti dan menyebut medan mana - bukan
menulis JSON yang separuh kosong dan membuat build gagal jauh di kemudian.
"""
import argparse
import json
import os
import re
import sys

import docx

KATA = {'zero': 0, 'one': 1, 'two': 2, 'three': 3, 'four': 4, 'five': 5,
        'six': 6, 'seven': 7, 'eight': 8, 'nine': 9, 'ten': 10, 'eleven': 11,
        'twelve': 12}
KELUARGA = {'Machine learning': 'ML', 'Traditional': 'trad'}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx', help='DOCX hasil build_paper.js sebelumnya')
    ap.add_argument('-o', '--outdir', default='.', help='dir-kerja tujuan')
    A = ap.parse_args()
    if not os.path.exists(A.docx):
        raise SystemExit(f'BERHENTI: {A.docx} tidak ada.')

    d = docx.Document(A.docx)
    teks = '\n'.join(p.text for p in d.paragraphs)

    # ---- rows: dari tabel satu langkah, yang sudah terurut menurut MASE.
    baris = []
    for t in d.tables:
        h = [c.text.strip() for c in t.rows[0].cells]
        if h[:2] == ['Method', 'Family'] and 'Best in' in h:
            for k, r in enumerate(t.rows[1:], 1):
                c = [x.text.strip() for x in r.cells]
                baris.append({
                    'model': c[0], 'family': KELUARGA[c[1]],
                    'mase': float(c[2]), 'med': float(c[3]),
                    'mae': float(c[4]), 'bias': float(c[5].replace('+', '')),
                    'best': int(c[6]), 'r1': k, 'r60': int(c[7])})
            break
    if not baris:
        raise SystemExit('BERHENTI: tabel satu langkah (Method | Family | ... | '
                         'Best in | Rank h=60) tidak ketemu di DOCX itu. '
                         'Yakin ini hasil build_paper.js, bukan naskah aslinya?')

    kurang = []

    def ambil(medan, pola, grup=1, ubah=float):
        m = re.search(pola, teks)
        if not m:
            kurang.append(medan)
            return None
        return ubah(m.group(grup))

    H1 = {
        'n_methods': len(baris),
        'n_train': ambil('n_train', r'on the first (\d+) business days', 1, int),
        'test_date': ambil('test_date',
                           r'test them on the final observation, ([^,]+),', 1, str),
        'n_series': ambil('n_series',
                          r'computed across the (\d+) terminal series', 1, int),
        'rec_tf_identical': ambil('rec_tf_identical',
                                  r'agree on all (\d+) of \d+ series-method pairs', 1, int),
        'rec_tf_pairs': ambil('rec_tf_pairs',
                              r'agree on all \d+ of (\d+) series-method pairs', 1, int),
        'dir_diff': ambil('dir_diff',
                          r'differs from the recursive path on (\d+) of the \d+ pairs', 1, int),
        'dir_pairs': ambil('dir_pairs',
                           r'differs from the recursive path on \d+ of the (\d+) pairs', 1, int),
        'dir_median_pct': ambil('dir_median_pct',
                                r'median absolute discrepancy of ([\d.]+) per cent'),
        'dir_max_pct': ambil('dir_max_pct', r'and a maximum of ([\d.]+) per cent'),
        'rows': baris,
    }

    # rec_tf_max dicetak sebagai "6 × 10⁻14"; dirakit kembali dari dua bagiannya.
    m = re.search(r'to within (\d+) × 10⁻(\d+)', teks)
    if m:
        H1['rec_tf_max'] = float(m.group(1)) * 10 ** -int(m.group(2))
    else:
        kurang.append('rec_tf_max')

    m = re.search(r'mean MASE of ([\d.]+) under the direct path against '
                  r'([\d.]+) under the other two', teks)
    if m:
        H1['proto_mase'] = {'Direct': float(m.group(1)),
                            'Recursive': float(m.group(2))}
    else:
        kurang.append('proto_mase')

    # rho di sini adalah korelasi peringkat SATU LANGKAH lawan ENAM PULUH
    # langkah - bukan rho Tahap I yang juga muncul di naskah sebagai -1,000.
    # Polanya sengaja menyebut "two columns of Table" supaya tidak tertukar.
    m = re.search(r'rank correlation between the two columns of Table \d+ is '
                  r'(-?[\d.]+) \(p = ([\d.]+)\)', teks)
    if m:
        H1['rho'], H1['rho_p'] = float(m.group(1)), float(m.group(2))
    else:
        kurang.extend(['rho', 'rho_p'])

    for medan, pola in (('n_below_one_mean', r'below unity — (\w+) of \w+ do so'),
                        ('n_below_one_med',
                         r'although (\w+) of \w+ achieve a median below unity')):
        m = re.search(pola, teks)
        if m and m.group(1) in KATA:
            H1[medan] = KATA[m.group(1)]
        else:
            kurang.append(medan)

    if kurang:
        raise SystemExit(
            'BERHENTI: medan berikut tidak ketemu di DOCX itu:\n'
            + ''.join(f'  - {k}\n' for k in kurang)
            + 'Kalimat yang memuatnya mungkin sudah diubah di build_paper.js. '
              'Perbarui polanya di skrip ini, jangan menulis JSON yang separuh '
              'kosong.')

    os.makedirs(A.outdir, exist_ok=True)
    keluar = os.path.join(A.outdir, 'h1_facts2.json')
    json.dump(H1, open(keluar, 'w'), indent=1)
    print(f'  {len(baris)} metode, {H1["n_series"]} seri, uji {H1["test_date"]}')
    print(f'  rho={H1["rho"]} (p={H1["rho_p"]}), '
          f'direct {H1["proto_mase"]["Direct"]} vs recursive '
          f'{H1["proto_mase"]["Recursive"]}')
    print(f'  ditulis {keluar}')
    print('\nCATATAN: nilainya sudah dibulatkan saat dicetak. Cukup untuk '
          'menyusun ulang\nnaskah, BUKAN pengganti menjalankan run_one_step.py.')


if __name__ == '__main__':
    main()
